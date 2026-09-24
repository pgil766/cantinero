"""Genera los documentos PDF y DOCX de la base de conocimiento y el PDF de la demo.

Salidas:
  data/seed/guia_tecnicas_de_bar.pdf
  data/seed/glosario_cocteleria.docx
  data/demo/guia_destilados_colombianos_viche.pdf   (NO va al seed)

Requiere fpdf2, python-docx y pypdf, y la fuente DejaVu Sans. En Windows (Smart App Control puede
bloquear lxml/Pillow) se ejecuta dentro de Docker, desde la raíz del repositorio:

    docker run --rm -v "%CD%/data:/work/data" python:3.12.14-slim bash -c \
      "apt-get update -qq && apt-get install -y -qq fonts-dejavu-core >/dev/null && \
       pip install -q fpdf2 python-docx pypdf && python /work/data/tools/generar_documentos.py"

Al final verifica que ningún archivo de data/seed/ mencione el viche (incluido el texto de PDF y DOCX).
"""
import datetime
import glob
import os
import re
import sys

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt
from fpdf import FPDF
from pypdf import PdfReader

sys.path.insert(0, os.path.dirname(__file__))
from contenido_documentos import DEMO_VICHE, GLOSARIO, GUIA_TECNICAS, WIKI_LICENCIA  # noqa: E402

DATA = os.path.abspath(os.path.join(os.path.dirname(__file__), ".."))
SEED = os.path.join(DATA, "seed")
DEMO = os.path.join(DATA, "demo")
FUENTE = "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf"
FUENTE_NEGRITA = "/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf"
HOY = datetime.date.today().isoformat()


class PDF(FPDF):
    def __init__(self, pie: str):
        super().__init__(format="A4")
        self.pie = pie
        self.add_font("DejaVu", "", FUENTE)
        self.add_font("DejaVu", "B", FUENTE_NEGRITA)
        self.set_auto_page_break(True, margin=18)
        self.set_margins(20, 18, 20)

    def footer(self):
        self.set_y(-12)
        self.set_font("DejaVu", "", 8)
        self.set_text_color(110)
        self.cell(0, 6, f"{self.pie} · página {self.page_no()}", align="C")
        self.set_text_color(0)


def pdf_documento(doc: dict, destino: str, pie: str, nota_licencia: str) -> None:
    pdf = PDF(pie)
    pdf.set_title(doc["titulo"])
    pdf.set_author("Proyecto Cantinero")
    pdf.add_page()
    pdf.set_font("DejaVu", "B", 20)
    pdf.multi_cell(0, 10, doc["titulo"], new_x="LMARGIN", new_y="NEXT")
    pdf.set_font("DejaVu", "", 12)
    pdf.set_text_color(80)
    pdf.multi_cell(0, 7, doc["subtitulo"], new_x="LMARGIN", new_y="NEXT")
    pdf.set_text_color(0)
    pdf.ln(3)
    pdf.set_font("DejaVu", "", 10.5)
    pdf.multi_cell(0, 5.8, doc["intro"], new_x="LMARGIN", new_y="NEXT")
    for titulo_seccion, bloques in doc["secciones"]:
        pdf.ln(4)
        pdf.set_font("DejaVu", "B", 14)
        pdf.multi_cell(0, 8, titulo_seccion, new_x="LMARGIN", new_y="NEXT")
        for subtitulo, texto in bloques:
            pdf.ln(1.5)
            if subtitulo:
                pdf.set_font("DejaVu", "B", 11)
                pdf.multi_cell(0, 6, subtitulo, new_x="LMARGIN", new_y="NEXT")
            pdf.set_font("DejaVu", "", 10.5)
            pdf.multi_cell(0, 5.8, texto, new_x="LMARGIN", new_y="NEXT")
    pdf.ln(5)
    pdf.set_font("DejaVu", "B", 14)
    pdf.multi_cell(0, 8, "Fuentes", new_x="LMARGIN", new_y="NEXT")
    pdf.set_font("DejaVu", "", 9.5)
    for f in doc["fuentes"]:
        pdf.multi_cell(0, 5.2, "• " + f, new_x="LMARGIN", new_y="NEXT")
    pdf.ln(3)
    pdf.set_font("DejaVu", "", 9)
    pdf.set_text_color(80)
    pdf.multi_cell(0, 5, f"{nota_licencia} Fecha de elaboración: {HOY}.", new_x="LMARGIN", new_y="NEXT")
    pdf.output(destino)


def docx_glosario(doc: dict, destino: str) -> None:
    d = Document()
    estilo = d.styles["Normal"]
    estilo.font.name = "Calibri"
    estilo.font.size = Pt(11)
    d.core_properties.title = doc["titulo"]
    d.core_properties.author = "Proyecto Cantinero"
    d.add_heading(doc["titulo"], level=0)
    d.add_paragraph(doc["intro"])
    letra = None
    for termino, definicion in sorted(doc["terminos"], key=lambda t: t[0].lower()):
        inicial = termino[0].upper()
        if inicial != letra:
            letra = inicial
            d.add_heading(letra, level=1)
        p = d.add_paragraph()
        p.add_run(termino + ": ").bold = True
        p.add_run(definicion)
    d.add_heading("Fuentes", level=1)
    for f in doc["fuentes"]:
        d.add_paragraph(f, style="List Bullet")
    nota = d.add_paragraph(f"{WIKI_LICENCIA} Fecha de elaboración: {HOY}.")
    nota.alignment = WD_ALIGN_PARAGRAPH.LEFT
    nota.runs[0].font.size = Pt(9)
    d.save(destino)


def texto_de(ruta: str) -> str:
    if ruta.endswith(".pdf"):
        return "\n".join(p.extract_text() or "" for p in PdfReader(ruta).pages)
    if ruta.endswith(".docx"):
        return "\n".join(p.text for p in Document(ruta).paragraphs)
    with open(ruta, encoding="utf-8") as f:
        return f.read()


def verificar_seed() -> int:
    patron = re.compile(r"[vb]iche", re.I)
    coincidencias = 0
    for ruta in sorted(glob.glob(os.path.join(SEED, "*"))):
        n = len(patron.findall(texto_de(ruta)))
        coincidencias += n
        if n:
            print(f"  ¡ATENCIÓN! {os.path.basename(ruta)}: {n} coincidencias de viche/biche")
    return coincidencias


def main() -> None:
    os.makedirs(SEED, exist_ok=True)
    os.makedirs(DEMO, exist_ok=True)
    guia = os.path.join(SEED, GUIA_TECNICAS["archivo"])
    pdf_documento(GUIA_TECNICAS, guia, "Cantinero · Guía de técnicas de bar", WIKI_LICENCIA)
    glos = os.path.join(SEED, GLOSARIO["archivo"])
    docx_glosario(GLOSARIO, glos)
    demo = os.path.join(DEMO, DEMO_VICHE["archivo"])
    pdf_documento(DEMO_VICHE, demo, "Cantinero · Guía de destilados colombianos", DEMO_VICHE["nota_licencia"])
    for ruta in (guia, glos, demo):
        texto = texto_de(ruta)
        print(f"{os.path.relpath(ruta, DATA)}: {os.path.getsize(ruta)} bytes, {len(texto)} caracteres de texto extraíble")
    total = verificar_seed()
    print(f"Coincidencias de viche/biche en data/seed/: {total}")
    if total:
        sys.exit(1)


if __name__ == "__main__":
    main()
